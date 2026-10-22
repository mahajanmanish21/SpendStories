"""Read credit card and bank statements (CSV, Excel, PDF) into Transactions.

Banks lay their exports out in wildly different ways, so nothing here is
bank-specific. Instead we look for a header row (Date / Description /
Amount or Debit+Credit columns), fall back to guessing columns from their
contents, and for PDFs fall back again to reading text lines that start
with a date and end with an amount.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from pathlib import Path

from .models import LoadResult, Transaction
from .parsing import DATE_TOKEN, detect_dayfirst, parse_amount, parse_date

SUPPORTED_EXTENSIONS = {".csv", ".tsv", ".txt", ".xlsx", ".xlsm", ".xls", ".pdf"}

# Column-name aliases, most specific first. Matching is on a normalised header.
ROLE_ALIASES: dict[str, list[str]] = {
    "balance": ["balance", "closing balance", "running balance", "available balance"],
    "date": [
        "transaction date", "txn date", "trans date", "tran date", "date",
        "posting date", "post date", "posted date", "value date", "booking date",
    ],
    "debit": [
        "debit", "debits", "withdrawal", "withdrawals", "withdrawal amt", "dr", "dr amount",
        "debit amount", "money out", "paid out", "spent", "outflow",
    ],
    "credit": [
        "credit", "credits", "deposit", "deposits", "deposit amt", "cr", "cr amount",
        "credit amount", "money in", "paid in", "received", "inflow",
    ],
    "crdr": ["dr/cr", "cr/dr", "dr / cr", "debit/credit", "credit/debit", "type", "txn type",
             "transaction type"],
    "amount": ["amount", "transaction amount", "txn amount", "amt", "billing amount",
               "amount in inr", "amount inr", "value"],
    "description": [
        "description", "narration", "transaction details", "transaction description",
        "details", "particulars", "merchant", "merchant name", "payee", "remarks",
        "memo", "name", "transaction remarks", "transaction",
    ],
}
# Roles are claimed in this order so e.g. "withdrawal amount" becomes debit, not amount.
ROLE_ORDER = ["balance", "date", "debit", "credit", "crdr", "amount", "description"]


def _norm_header(cell) -> str:
    s = str(cell or "").lower()
    s = re.sub(r"\(.*?\)", " ", s)          # "Amount (INR)" -> "amount"
    s = re.sub(r"[^a-z/ ]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def _match_role(header: str) -> str | None:
    if not header or len(header.split()) > 5:
        return None
    for role in ROLE_ORDER:
        for alias in ROLE_ALIASES[role]:
            if header == alias or header.startswith(alias + " ") or header.endswith(" " + alias):
                return role
    return None


@dataclass
class ColumnMap:
    date: int
    description: int
    amount: int | None = None
    debit: int | None = None
    credit: int | None = None
    crdr: int | None = None
    balance: int | None = None


def find_header(rows: list[list[str]], scan: int = 80) -> tuple[int, ColumnMap] | None:
    for i, row in enumerate(rows[:scan]):
        roles: dict[str, int] = {}
        for j, cell in enumerate(row):
            role = _match_role(_norm_header(cell))
            if role and role not in roles:
                roles[role] = j
        if "date" in roles and "description" in roles and (
            "amount" in roles or "debit" in roles or "credit" in roles
        ):
            return i, ColumnMap(**roles)
    return None


def _guess_columns(rows: list[list[str]]) -> ColumnMap | None:
    """Header-less file: infer columns from what the cells contain."""
    sample = [r for r in rows if any(str(c).strip() for c in r)][:200]
    if not sample:
        return None
    width = max(len(r) for r in sample)
    date_hits = [0] * width
    num_hits = [0] * width
    text_len = [0] * width
    for r in sample:
        for j in range(width):
            c = str(r[j]).strip() if j < len(r) else ""
            if not c:
                continue
            if parse_date(c):
                date_hits[j] += 1
            elif parse_amount(c)[0] is not None:
                num_hits[j] += 1
            else:
                text_len[j] += len(c)
    date_col = max(range(width), key=lambda j: date_hits[j])
    if date_hits[date_col] < len(sample) * 0.5:
        return None
    desc_col = max((j for j in range(width) if j != date_col), key=lambda j: text_len[j], default=None)
    nums = [j for j in range(width) if num_hits[j] >= len(sample) * 0.3 and j not in (date_col, desc_col)]
    if desc_col is None or not nums:
        return None
    # A single numeric column is the amount; with several, the first is the amount
    # and a trailing one that is always filled is most likely a running balance.
    return ColumnMap(date=date_col, description=desc_col, amount=nums[0])


_SUMMARY_ROW = re.compile(
    r"opening balance|closing balance|balance b/?f|balance forward|brought forward|carried forward"
    r"|^total\b|statement summary|page \d", re.I)
_CREDIT_WORDS = ("credit", "deposit", "refund", "payment", "reversal", "cashback")
_DEBIT_WORDS = ("debit", "withdrawal", "purchase", "sale", "charge", "fee")


def _cell(row: list, idx: int | None):
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def extract_from_rows(rows: list[list], source: str, sign: str = "auto",
                      dayfirst: bool | None = None) -> tuple[list[Transaction], list[str]]:
    warnings: list[str] = []
    header = find_header(rows)
    if header:
        start, cols = header[0] + 1, header[1]
    else:
        cols = _guess_columns(rows)
        start = 0
        if cols is None:
            return [], ["could not find Date/Description/Amount columns"]
        warnings.append("no header row found; columns were guessed from their contents")

    body = rows[start:]
    if dayfirst is None:
        dayfirst = detect_dayfirst(_cell(r, cols.date) for r in body)

    parsed = []  # (date, desc, value, marker, explicit_sign)
    last = None
    for r in body:
        d = parse_date(_cell(r, cols.date), dayfirst)
        desc = re.sub(r"\s+", " ", str(_cell(r, cols.description) or "")).strip()
        if d is None:
            # A wrapped description line (no date, no numbers) continues the previous row.
            numeric_cols = (cols.amount, cols.debit, cols.credit, cols.balance)
            if (last is not None and desc and len(last[1]) < 200 and not _SUMMARY_ROW.search(desc)
                    and not any(parse_amount(_cell(r, c))[0] for c in numeric_cols if c is not None)):
                last[1] = f"{last[1]} {desc}"
            continue
        if _SUMMARY_ROW.search(desc):
            continue

        if cols.debit is not None or cols.credit is not None:
            debit = parse_amount(_cell(r, cols.debit))[0] or 0.0
            credit = parse_amount(_cell(r, cols.credit))[0] or 0.0
            value = abs(debit) - abs(credit)
            if value == 0 and cols.amount is not None:
                value, marker = parse_amount(_cell(r, cols.amount))
                if value is None:
                    continue
                last = [d, desc, value, marker, False]
                parsed.append(last)
                continue
            if debit == 0 and credit == 0:
                continue
            last = [d, desc, value, None, True]
        else:
            value, marker = parse_amount(_cell(r, cols.amount))
            if value is None:
                continue
            crdr = str(_cell(r, cols.crdr) or "").strip().lower()
            if crdr in ("c", "cr", "cr.") or crdr.startswith(_CREDIT_WORDS):
                marker = "cr"
            elif crdr in ("d", "dr", "dr.") or crdr.startswith(_DEBIT_WORDS):
                marker = "dr"
            last = [d, desc, value, marker, False]
        parsed.append(last)

    return _finalise(parsed, source, sign, warnings), warnings


def _finalise(parsed, source, sign, warnings) -> list[Transaction]:
    # For a single signed amount column we need to know whether spends are
    # positive (most card exports) or negative (most bank exports). Since most
    # rows in any statement are spends, the majority sign tells us.
    unmarked = [p[2] for p in parsed if not p[4] and p[3] is None and p[2] != 0]
    if sign == "auto":
        negatives = sum(1 for v in unmarked if v < 0)
        spend_negative = negatives > len(unmarked) / 2
        if unmarked and spend_negative:
            warnings.append("amounts look like 'negative = spend'; signs were flipped")
    else:
        spend_negative = sign == "negative-spend"

    txns = []
    for d, desc, value, marker, explicit in parsed:
        if explicit:
            amount = value
        elif marker == "cr":
            amount = -abs(value)
        elif marker == "dr":
            amount = abs(value)
        else:
            amount = -value if spend_negative else value
        if amount == 0:
            continue
        txns.append(Transaction(date=d, description=desc or "(no description)",
                                amount=round(amount, 2), source=source))
    return txns


# --------------------------------------------------------------------------- CSV / Excel

def _read_csv_rows(path: Path) -> list[list[str]]:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    sample = text[:20000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    return [row for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def _read_excel_sheets(path: Path) -> list[list[list]]:
    if path.suffix.lower() == ".xls":
        import pandas as pd  # needs xlrd for legacy .xls

        try:
            sheets = pd.read_excel(path, sheet_name=None, header=None, dtype=object)
        except ImportError as e:
            raise RuntimeError("legacy .xls needs the 'xlrd' package (pip install xlrd), "
                               "or re-save the file as .xlsx/.csv") from e
        return [df.where(df.notna(), None).values.tolist() for df in sheets.values()]

    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        sheets.append([["" if c is None else c for c in row] for row in ws.iter_rows(values_only=True)])
    wb.close()
    return sheets


# --------------------------------------------------------------------------- PDF

_MONEY = re.compile(r"(?<![\w.,])(-?\(?[\d,]*\d\.\d{2}\)?-?)(\s*(?:cr|dr)\b\.?)?", re.I)
_LEADING_DATE = re.compile(rf"^\s*({DATE_TOKEN.pattern})\s+(.*)$")


def _pdf_text_lines(pages) -> list[tuple[str, str, list[tuple[float, str | None]]]]:
    """Pull out lines shaped like '<date> <description> <amount> [<amount>...]'."""
    out = []
    for page in pages:
        text = page.extract_text() or ""
        for line in text.splitlines():
            m = _LEADING_DATE.match(line)
            if not m:
                continue
            date_str, rest = m.group(1), m.group(2)
            # Some statements print a second (value/posting) date right after the first.
            m2 = _LEADING_DATE.match(rest)
            if m2 and parse_date(m2.group(1)):
                rest = m2.group(2)
            amounts = []
            # Collect money tokens from the end of the line backwards.
            tail = rest
            while True:
                mm = None
                for mm in _MONEY.finditer(tail):
                    pass
                if mm is None or tail[mm.end():].strip():
                    break
                value, _ = parse_amount(mm.group(1))
                marker = (mm.group(2) or "").strip().lower().rstrip(".") or None
                amounts.insert(0, (value, marker))
                tail = tail[: mm.start()].rstrip()
            if not amounts or not tail.strip():
                continue
            out.append((date_str, tail.strip(), amounts))
    return out


def _extract_pdf_text(pages, source, sign) -> tuple[list[Transaction], list[str]]:
    lines = _pdf_text_lines(pages)
    dayfirst = detect_dayfirst(l[0] for l in lines)
    parsed = []
    prev_balance = None
    for date_str, desc, amounts in lines:
        d = parse_date(date_str, dayfirst)
        if d is None or _SUMMARY_ROW.search(desc):
            if len(amounts) == 1:
                prev_balance = amounts[0][0]
            continue
        if len(amounts) == 1:
            value, marker = amounts[0]
            parsed.append([d, desc, value, marker, False])
            continue
        # Bank layout: [..., amount, balance] or [debit, credit, balance].
        balance = amounts[-1][0]
        if len(amounts) >= 3:
            debit, credit = abs(amounts[-3][0]), abs(amounts[-2][0])
            parsed.append([d, desc, debit - credit, None, True])
        else:
            value, marker = amounts[-2]
            value = abs(value)
            if marker is None and prev_balance is not None and abs(abs(balance - prev_balance) - value) < 0.015:
                parsed.append([d, desc, value if balance < prev_balance else -value, None, True])
            else:
                parsed.append([d, desc, value, marker, False])
        prev_balance = balance
    warnings: list[str] = []
    return _finalise(parsed, source, sign, warnings), warnings


def _load_pdf(path: Path, source: str, sign: str, password: str | None) -> LoadResult:
    import pdfplumber

    result = LoadResult(source=source)
    try:
        pdf = pdfplumber.open(path, password=password or "")
    except Exception as e:  # pdfminer raises several types for bad passwords
        name = type(e).__name__
        if "Password" in name or "password" in str(e).lower():
            result.warnings.append("PDF is password protected; pass --pdf-password")
            return result
        raise
    with pdf:
        pages = list(pdf.pages)
        # 1) Ruled tables. Stitch tables across pages: the header usually only
        #    appears on the first page.
        rows: list[list] = []
        for page in pages:
            for table in page.extract_tables() or []:
                rows.extend([c if c is not None else "" for c in row] for row in table)
        table_txns, table_warn = (extract_from_rows(rows, source, sign) if rows else ([], []))
        # 2) Text lines.
        text_txns, text_warn = _extract_pdf_text(pages, source, sign)

    if len(table_txns) >= len(text_txns) and table_txns:
        result.transactions, result.method = table_txns, "pdf-table"
        result.warnings.extend(w for w in table_warn if "no header" not in w)
    else:
        result.transactions, result.method = text_txns, "pdf-text"
        result.warnings.extend(text_warn)
    if result.transactions:
        result.warnings.append("PDF parsing is best-effort; spot-check the transactions CSV "
                               "or use a CSV/Excel export if your bank offers one")
    return result


# --------------------------------------------------------------------------- entry point

def load_file(path: str | Path, sign: str = "auto", pdf_password: str | None = None,
              dayfirst: bool | None = None) -> LoadResult:
    path = Path(path)
    source = path.stem
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _load_pdf(path, source, sign, pdf_password)

    result = LoadResult(source=source)
    if ext in (".xlsx", ".xlsm", ".xls"):
        sheets = _read_excel_sheets(path)
        result.method = "excel"
    else:
        sheets = [_read_csv_rows(path)]
        result.method = "csv"
    for rows in sheets:
        txns, warns = extract_from_rows(rows, source, sign, dayfirst)
        if txns:
            result.transactions.extend(txns)
            result.warnings.extend(warns)
    if not result.transactions:
        result.warnings.append("no transactions found")
    return result


def collect_files(inputs: list[str | Path]) -> list[Path]:
    files: list[Path] = []
    for item in inputs:
        p = Path(item)
        if p.is_dir():
            files.extend(sorted(f for f in p.rglob("*")
                                if f.is_file() and f.suffix.lower() in SUPPORTED_EXTENSIONS
                                and not f.name.startswith((".", "~$"))))
        elif p.is_file():
            files.append(p)
        else:
            raise FileNotFoundError(item)
    return files
